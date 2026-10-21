import csv
import io
import os
import re
import uuid
from typing import Dict, List, Optional
from reportlab.lib.pagesizes import letter
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
import docx
from app.config import settings
from app.logging_config import get_logger

logger = get_logger(__name__)


class FileGenerationService:
    def __init__(self, base_dir: Optional[str] = None):
        self.base_dir = os.path.abspath(base_dir or os.path.join(settings.STORAGE_DIR, "temp"))
        os.makedirs(self.base_dir, exist_ok=True)

    def _sanitize_filename(self, filename: str, default_ext: str = ".txt") -> str:
        # Strip path traversal and forbidden characters
        base = os.path.basename(filename.replace("\\", "/"))
        clean = re.sub(r"[^\w\s\-\.]", "", base).strip()
        clean = re.sub(r"\.{2,}", "", clean)  # Remove any sequence of multiple dots
        if not clean or clean.startswith("."):
            clean = f"export_{uuid.uuid4().hex[:8]}"
        if not os.path.splitext(clean)[1]:
            clean += default_ext
        return clean

    def _get_safe_path(self, filename: str) -> str:
        safe_name = self._sanitize_filename(filename)
        unique_prefix = uuid.uuid4().hex[:6]
        full_path = os.path.abspath(os.path.join(self.base_dir, f"{unique_prefix}_{safe_name}"))
        
        # Prevent path traversal
        if not full_path.startswith(self.base_dir):
            raise ValueError("Path traversal attempt detected in filename.")
        return full_path

    def generate_text_file(self, content: str, filename: str = "notes.txt") -> str:
        path = self._get_safe_path(filename)
        with open(path, "w", encoding="utf-8") as f:
            f.write(content)
        return path

    def generate_csv_file(self, rows: List[Dict[str, str]], filename: str = "data.csv") -> str:
        if not rows:
            raise ValueError("No data provided to generate CSV.")
        path = self._get_safe_path(filename)
        fieldnames = list(rows[0].keys())
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)
        return path

    def generate_pdf_report(self, title: str, paragraphs: List[str], filename: str = "report.pdf") -> str:
        path = self._get_safe_path(filename)
        doc = SimpleDocTemplate(path, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            "TitleStyle",
            parent=styles["Heading1"],
            fontSize=18,
            leading=22,
            spaceAfter=15,
        )
        body_style = ParagraphStyle(
            "BodyStyle",
            parent=styles["Normal"],
            fontSize=11,
            leading=16,
            spaceAfter=10,
        )

        elements = [Paragraph(title, title_style), Spacer(1, 10)]
        for p_text in paragraphs:
            safe_text = p_text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            elements.append(Paragraph(safe_text, body_style))
            elements.append(Spacer(1, 6))

        doc.build(elements)
        return path

    def generate_docx_document(self, title: str, paragraphs: List[str], filename: str = "document.docx") -> str:
        path = self._get_safe_path(filename)
        doc = docx.Document()
        doc.add_heading(title, 0)
        for p in paragraphs:
            doc.add_paragraph(p)
        doc.save(path)
        return path


file_generator = FileGenerationService()
